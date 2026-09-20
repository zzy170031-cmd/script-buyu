"""Editable story-only DOCX and strict visible-content round trip."""
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from docx_io import package, visible_text, W
from story_contract import ContractError, digest, require, validate_story, write_bytes_new

MARKER = "script-buyu/screenplay/1.0"
META = [("project_id", "项目编号"), ("genre", "类型"), ("domain", "故事背景"),
        ("audience", "受众"), ("medium", "表现形式"), ("viewpoint", "叙述视角"),
        ("episode_count", "集数"), ("target_duration_seconds", "总时长预算 秒")]
SECTIONS = ["故事概览", "故事梗概", "人物", "世界规则", "分集剧本"]
CHAR_HEADERS = ["人物ID", "姓名", "人物设定", "人物弧线"]
CHAR_KEYS = ["character_id", "name", "description", "arc"]
STYLES = ["SBLogline", "SBSynopsis", "SBWorldRule", "SBEpisode", "SBEpisodeMeta",
          "SBScene", "SBLocation", "SBTime", "SBPurpose", "SBAction", "SBDialogue"]


def _style(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in STYLES:
        base = "Heading 1" if name == "SBEpisode" else "Heading 2" if name == "SBScene" else "Normal"
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles[base]
    for name in ["Normal", "Title", "Heading 1", "Heading 2"] + STYLES:
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["SBDialogue"].paragraph_format.left_indent = Inches(.2)
    for name in ["SBEpisode", "SBEpisodeMeta", "SBScene", "SBLocation", "SBTime", "SBPurpose"]:
        doc.styles[name].paragraph_format.keep_with_next = True


def _table(doc, rows, widths):
    table = doc.add_table(rows=0, cols=len(widths))
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                node = OxmlElement("w:" + side)
                node.set(qn("w:w"), "90")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            properties.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                node = OxmlElement("w:" + side)
                for key, val in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    node.set(qn("w:" + key), val)
                borders.append(node)
            properties.append(borders)
            if index == 0:
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), "E6EDF2")
                properties.append(fill)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            table.rows[index]._tr.get_or_add_trPr().append(repeat)
    return table


def _bookmark(paragraph, identifier, serial):
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(serial))
    start.set(qn("w:name"), "SBB_" + identifier)
    end.set(qn("w:id"), str(serial))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def create_docx(story, output):
    validate_story(story)
    doc = Document()
    _style(doc)
    doc.core_properties.identifier = MARKER
    doc.core_properties.title = story["project"]["title"]
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.add_paragraph(story["project"]["title"], "Title")
    doc.add_paragraph("故事剧本供内容确认，确认范围为人物、情节、对白与分集安排。", "Subtitle")
    rows = [["项目资料", "内容"], ["版本", story["revision"]], ["状态", "待用户确认"]]
    rows += [[label, story["project"][key]] for key, label in META]
    _table(doc, rows, [1.55, 5.22])
    for heading, key, style in [("故事概览", "logline", "SBLogline"), ("故事梗概", "synopsis", "SBSynopsis")]:
        doc.add_heading(heading, 1)
        doc.add_paragraph(story[key], style)
    doc.add_heading("人物", 1)
    _table(doc, [CHAR_HEADERS] + [[c[key] for key in CHAR_KEYS] for c in story["characters"]], [.62, .85, 2.9, 2.4])
    doc.add_heading("世界规则", 1)
    for rule in story["world_rules"]:
        doc.add_paragraph(rule, "SBWorldRule")
    doc.add_heading("分集剧本", 1)
    names = {c["character_id"]: c["name"] for c in story["characters"]}
    serial = 0
    for index, episode in enumerate(story["episodes"], 1):
        p = doc.add_paragraph(f"第{index}集 {episode['title']}", "SBEpisode")
        p.paragraph_format.page_break_before = True
        doc.add_paragraph(f"{episode['episode_id']}｜{episode['target_duration_seconds']}秒", "SBEpisodeMeta")
        for scene in [s for s in story["scenes"] if s["episode_id"] == episode["episode_id"]]:
            doc.add_paragraph(scene["scene_id"] + "｜" + scene["heading"], "SBScene")
            for field, style in [("location", "SBLocation"), ("time_weather", "SBTime"), ("purpose", "SBPurpose")]:
                doc.add_paragraph(scene[field], style)
            for beat in scene["beats"]:
                serial += 1
                prefix = names[beat["speaker_id"]] + "：" if beat["kind"] == "dialogue" else ""
                p = doc.add_paragraph(prefix + beat["text"], "SBDialogue" if prefix else "SBAction")
                _bookmark(p, beat["beat_id"], serial)
    data = BytesIO()
    doc.save(data)
    require(digest(extract_docx(data)) == digest(story), "DOCX round trip changed story data")
    write_bytes_new(output, data.getvalue())


def _check_supported(parts):
    forbidden = {W + tag for tag in ("ins", "del", "moveFrom", "moveTo", "drawing", "pict", "object", "txbxContent", "fldSimple", "fldChar", "sdt", "customXml", "hyperlink", "altChunk", "commentRangeStart", "vanish", "gridSpan", "vMerge")}
    for name, root in parts.items():
        if name.startswith("word/") and name.endswith(".xml") and name not in ("word/styles.xml", "word/settings.xml", "word/fontTable.xml", "word/numbering.xml", "word/webSettings.xml") and "/theme/" not in name:
            require(not any(node.tag in forbidden for node in root.iter()), "unsupported edits or objects; use full-source review before merging")
            if name != "word/document.xml":
                require(not visible_text(root).strip(), "non-body content requires explicit review")
    styles = parts.get("word/styles.xml")
    if styles is not None:
        by_id = {style.get(W + "styleId"): style for style in styles.findall(W + "style")}
        used = {node.get(W + "val") for node in parts["word/document.xml"].iter() if node.tag in (W + "pStyle", W + "rStyle", W + "tblStyle")}
        used.update(key for key, style in by_id.items() if style.get(W + "default") == "1")
        defaults = styles.find(W + "docDefaults")
        inspect = [defaults] if defaults is not None else []
        visited = set()
        while used:
            key = used.pop()
            if key in visited or key not in by_id:
                continue
            visited.add(key)
            style = by_id[key]
            inspect.append(style)
            used.update(node.get(W + "val") for node in style if node.tag in (W + "basedOn", W + "link"))
        for element in inspect:
            require(not any(node.tag in (W + "vanish", W + "webHidden") and node.get(W + "val", "1") not in ("0", "false", "off") for node in element.iter()), "referenced style hides text; explicit review required")


def _check_paragraph(p):
    allowed = {qn("w:" + tag) for tag in ("pPr", "r", "bookmarkStart", "bookmarkEnd", "proofErr")}
    runs = {qn("w:" + tag) for tag in ("rPr", "t", "tab", "br", "cr", "lastRenderedPageBreak")}
    require(all(child.tag in allowed for child in p._p), "unsupported paragraph container; review before merging")
    require(all(child.tag in runs for run in p.runs for child in run._r), "unsupported run content; review before merging")
    require(p.text == visible_text(p._p), "paragraph contains text outside normal runs")


def extract_docx(source):
    raw, parts = package(source)
    _check_supported(parts)
    doc = Document(BytesIO(raw))
    require(doc.core_properties.identifier == MARKER, "not a script-buyu screenplay; use read_source_docx.py")
    story = {"schema_version": "1.0", "status": "awaiting_user_confirmation", "project": {}, "characters": [], "world_rules": [], "episodes": [], "scenes": []}
    section_index, tables = -1, 0
    scene, episode, pending_episode = None, None, None
    seen = set()
    for node in doc.element.body:
        if node.tag == qn("w:sectPr"):
            continue
        if node.tag == qn("w:tbl"):
            table = Table(node, doc)
            require(all(child.tag in (qn("w:tblPr"), qn("w:tblGrid"), qn("w:tr")) for child in node), "unsupported table-level structure; review before merging")
            for row in table.rows:
                require(all(child.tag in (qn("w:trPr"), qn("w:tc")) for child in row._tr), "unsupported row structure; review before merging")
            require(all(len(row.cells) == len(table.columns) for row in table.rows), "irregular table")
            require(not node.xpath(".//w:tbl"), "nested table")
            for row in table.rows:
                for cell in row.cells:
                    require(all(child.tag in (qn("w:tcPr"), qn("w:p")) for child in cell._tc), "unsupported table-cell content")
                    for paragraph in cell.paragraphs:
                        _check_paragraph(paragraph)
            values = [[cell.text for cell in row.cells] for row in table.rows]
            require(bool(values), "empty table")
            if tables == 0 and section_index == -1:
                require(len(values) == len(META) + 3 and values[0] == ["项目资料", "内容"], "project table changed")
                expected = ["版本", "状态"] + [label for _, label in META]
                require([r[0] for r in values[1:]] == expected, "project fields missing or reordered")
                require(values[2][1] == "待用户确认", "document cannot record or imply approval")
                try:
                    story["revision"] = int(values[1][1])
                    for (key, _), row in zip(META, values[3:]):
                        value = row[1]
                        story["project"][key] = int(value) if key == "episode_count" else float(value) if key == "target_duration_seconds" else value
                except ValueError as exc:
                    raise ContractError("invalid project numeric field") from exc
            elif tables == 1 and section_index == 2:
                require(values[0] == CHAR_HEADERS and all(len(row) == 4 for row in values), "character table changed")
                story["characters"] = [dict(zip(CHAR_KEYS, row)) for row in values[1:]]
            else:
                raise ContractError("unexpected table; content cannot be silently omitted")
            tables += 1
            continue
        require(node.tag == qn("w:p"), "unsupported body element")
        p = Paragraph(node, doc)
        text, style = p.text, p.style.name
        _check_paragraph(p)
        if not text.strip():
            continue
        if style == "Title" and section_index == -1 and "title" not in story["project"]:
            story["project"]["title"] = text
        elif style == "Subtitle" and section_index == -1 and "subtitle" not in seen:
            require(text == "故事剧本供内容确认，确认范围为人物、情节、对白与分集安排。", "unexpected introductory text")
            seen.add("subtitle")
        elif style == "Heading 1":
            section_index += 1
            require(section_index < len(SECTIONS) and text == SECTIONS[section_index], "section missing, duplicated or changed")
        elif style in ("SBLogline", "SBSynopsis"):
            key, expected = ("logline", 0) if style == "SBLogline" else ("synopsis", 1)
            require(section_index == expected and key not in story, "duplicate or misplaced story summary")
            story[key] = text
        elif style == "SBWorldRule" and section_index == 3:
            story["world_rules"].append(text)
        elif style == "SBEpisode" and section_index == 4:
            require(pending_episode is None, "missing episode metadata")
            match = re.fullmatch(r"第(\d+)集 (.+)", text, re.S)
            require(match is not None and int(match[1]) == len(story["episodes"]) + 1, "invalid episode heading")
            pending_episode = match[2]
            scene = None
        elif style == "SBEpisodeMeta" and pending_episode is not None:
            match = re.fullmatch(r"([A-Za-z0-9_.-]+)｜([0-9.eE+\-]+)秒", text)
            require(match is not None, "invalid episode metadata")
            try:
                seconds = float(match[2])
            except ValueError as exc:
                raise ContractError("invalid duration") from exc
            episode = {"episode_id": match[1], "title": pending_episode, "target_duration_seconds": seconds}
            story["episodes"].append(episode)
            pending_episode = None
        elif style == "SBScene" and section_index == 4 and episode is not None and pending_episode is None:
            fields = text.split("｜", 1)
            require(len(fields) == 2, "invalid scene heading")
            scene = {"scene_id": fields[0], "episode_id": episode["episode_id"], "heading": fields[1], "beats": []}
            story["scenes"].append(scene)
        elif style in ("SBLocation", "SBTime", "SBPurpose") and scene is not None:
            key = {"SBLocation": "location", "SBTime": "time_weather", "SBPurpose": "purpose"}[style]
            require(key not in scene, "duplicate scene field")
            scene[key] = text
        elif style in ("SBAction", "SBDialogue") and scene is not None:
            anchors = [n.get(qn("w:name")) for n in node.xpath(".//w:bookmarkStart") if n.get(qn("w:name"), "").startswith("SBB_")]
            require(len(anchors) == 1, "missing or duplicate story beat bookmark; review this edit")
            beat = {"beat_id": anchors[0][4:], "kind": "action", "text": text, "speaker_id": None}
            if style == "SBDialogue":
                fields = text.split("：", 1)
                by_name = {c["name"]: c["character_id"] for c in story["characters"]}
                require(len(fields) == 2 and fields[0] in by_name, "unknown dialogue speaker")
                beat.update(kind="dialogue", text=fields[1], speaker_id=by_name[fields[0]])
            scene["beats"].append(beat)
        else:
            raise ContractError("unrecognized paragraph or position: " + style + "; use full-source review")
    require(tables == 2 and section_index == 4 and pending_episode is None, "incomplete screenplay document")
    return validate_story(story)
