"""Embed local character portraits and optional prompts into a new screenplay DOCX."""
import argparse
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.image.image import Image
from docx.oxml.ns import qn
from docx.shared import Inches

from screenplay_docx import CHAR_HEADERS, extract_docx
from story_contract import ContractError, read_json, require, write_bytes_new


def embed_images(source, mapping, output, image_base):
    source=Path(source)
    output=Path(output)
    require(source.resolve()!=output.resolve(),'image edition requires a new output path')
    require(not output.exists(),'output already exists')
    story=extract_docx(source)
    require(isinstance(mapping,dict) and bool(mapping),'images must be a nonempty character-id mapping')
    ids={c['character_id']:c['name'] for c in story['characters']}
    require(set(mapping)<=set(ids),'image mapping contains an unknown character id')
    validated={}
    prompts={}
    for character_id,value in mapping.items():
        if isinstance(value,dict):
            require(set(value)=={'prompt','image'},'visual entry requires prompt and image')
            require(isinstance(value['prompt'],str) and value['prompt'].strip(),'portrait prompt must be nonempty')
            prompts[character_id]=value['prompt']
            value=value['image']
        require(isinstance(value,str) and value.strip(),'image path must be a string')
        require('://' not in value and not value.startswith(('\\\\','//')),'only local images are supported')
        p=Path(value)
        p=(Path(image_base)/p).resolve() if not p.is_absolute() else p.resolve()
        require(p.is_file() and p.suffix.lower() in ('.png','.jpg','.jpeg'),'expected a local PNG or JPEG')
        require(p.stat().st_size<=20*1024*1024,'image exceeds 20 MiB')
        blob=p.read_bytes()
        try:
            info=Image.from_file(BytesIO(blob))
        except Exception as exc:
            raise ContractError('image is not a readable PNG or JPEG') from exc
        require(info.content_type in ('image/png','image/jpeg'),'unsupported image encoding')
        require(info.px_width>0 and info.px_height>0,'invalid image dimensions')
        validated[character_id]=(blob,info)
    doc=Document(source)
    table=doc.tables[1]
    require([c.text for c in table.rows[0].cells]==CHAR_HEADERS,'character table differs from supported format')
    table.add_column(Inches(1.2))
    table.rows[0].cells[-1].text='人物形象提示词与参考图' if prompts else '人物图片'
    for row in table.rows:
        target=row.cells[-1]
        source_properties=row.cells[-2]._tc.get_or_add_tcPr()
        target._tc.remove(target._tc.get_or_add_tcPr())
        target._tc.insert(0,deepcopy(source_properties))
    for run in table.rows[0].cells[-1].paragraphs[0].runs: run.bold=True
    widths=[.45,.6,1.8,1.35,2.57] if prompts else [.6,.8,2.15,2.02,1.2]
    for col,width in zip(table.columns,widths): col.width=Inches(width)
    for row in table.rows:
        for cell,width in zip(row.cells,widths): cell.width=Inches(width)
    for row in table.rows[1:]:
        character_id=row.cells[0].text
        if character_id not in validated: continue
        blob,info=validated[character_id]
        cell=row.cells[-1]
        if character_id in prompts:
            cell.paragraphs[0].text=prompts[character_id]
            image_p=cell.add_paragraph()
        else:
            image_p=cell.paragraphs[0]
        width=min(1.03,1.4*info.px_width/info.px_height)
        if prompts: width=min(2.1,2.6*info.px_width/info.px_height)
        picture=image_p.add_run().add_picture(BytesIO(blob),width=Inches(width))
        picture._inline.docPr.set('descr',ids[character_id]+' 人物参考图片')
    data=BytesIO();doc.save(data)
    reopened=Document(BytesIO(data.getvalue()))
    require(len(reopened.inline_shapes)==len(mapping),'embedded image count mismatch')
    require([[c.text for c in r.cells[:4]] for r in reopened.tables[1].rows]==
            [CHAR_HEADERS]+[[c[k] for k in ('character_id','name','description','arc')] for c in story['characters']],
            'portrait insertion changed character text')
    require([p.text for p in reopened.paragraphs]==[p.text for p in Document(source).paragraphs],
            'portrait insertion changed story text')
    for row in reopened.tables[1].rows[1:]:
        if row.cells[0].text in prompts:
            require(row.cells[-1].paragraphs[0].text==prompts[row.cells[0].text], 'portrait prompt changed')
    write_bytes_new(output,data.getvalue())


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input',required=True)
    p.add_argument('--images',required=True)
    p.add_argument('--output',required=True)
    a=p.parse_args()
    try:
        embed_images(a.input,read_json(a.images),a.output,Path(a.images).resolve().parent)
    except (ContractError,OSError,ValueError) as exc:
        p.exit(2,'Cannot embed portraits: '+str(exc)+'\n')
    print('Portrait edition created. Verify identity and rendered layout. Use full-source review for later edits.')


if __name__=='__main__': main()
