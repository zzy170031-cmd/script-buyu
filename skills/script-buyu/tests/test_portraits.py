import base64
import tempfile
import unittest
import sys
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
from docx import Document

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))

from fixtures import story
from screenplay_docx import create_docx,extract_docx
from embed_character_images import embed_images
from story_contract import ContractError

# Original one-pixel fixture, no user character or real likeness.
PNG=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aZ1sAAAAASUVORK5CYII=')


class PortraitTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.source=self.root/'story.docx';create_docx(story(),self.source)
        (self.root/'portrait.png').write_bytes(PNG)

    def test_embedding_keeps_text_and_packages_picture(self):
        out=self.root/'pictured.docx'
        embed_images(self.source,{'C1':'portrait.png'},out,self.root)
        doc=Document(out)
        self.assertEqual(len(doc.inline_shapes),1)
        self.assertEqual([p.text for p in doc.paragraphs],[p.text for p in Document(self.source).paragraphs])
        self.assertEqual(doc.tables[1].rows[1].cells[1].text,'阿砚')
        with ZipFile(out) as z: self.assertTrue(any(n.startswith('word/media/') for n in z.namelist()))
        with self.assertRaises(ContractError): extract_docx(out)

    def test_unknown_character_remote_and_overwrite_fail(self):
        for mapping,out in [({'C999':'portrait.png'},self.root/'bad.docx'),
                            ({'C1':'https://example.com/image.png'},self.root/'remote.docx'),
                            ({'C1':'portrait.png'},self.source)]:
            with self.assertRaises(ContractError): embed_images(self.source,mapping,out,self.root)
        self.assertEqual(extract_docx(self.source),story())

    def test_prompt_and_picture_remain_paired(self):
        out=self.root/'designed.docx'
        prompt='阿砚，短发，素色外套，符合当前故事的水彩人物参考图。'
        embed_images(self.source,{'C1':{'prompt':prompt,'image':'portrait.png'}},out,self.root)
        doc=Document(out)
        self.assertEqual(doc.tables[1].rows[0].cells[-1].text,'人物形象提示词与参考图')
        self.assertEqual(doc.tables[1].rows[1].cells[-1].paragraphs[0].text,prompt)
        self.assertEqual(len(doc.inline_shapes),1)
        self.assertEqual([p.text for p in doc.paragraphs],[p.text for p in Document(self.source).paragraphs])
        for value in ({'prompt':'','image':'portrait.png'}, {'prompt':prompt,'image':'https://example.com/a.png'}):
            with self.assertRaises(ContractError):
                embed_images(self.source,{'C1':value},self.root/'invalid.docx',self.root)

    def test_no_person_silent_story_roundtrip(self):
        data=story();data['characters']=[]
        data['project'].update(title='空屋雨声',genre='观察短片',domain='无人旧屋')
        data['logline']='一场雨从空屋的破窗进入，雨止后积水映出天空。'
        data['synopsis']='风推开破窗，雨水落在地板上。雨渐歇，水面平静下来，映出檐口后的天光。'
        data['world_rules']=['房屋内没有人物，不使用对白或旁白。']
        for episode,title in zip(data['episodes'],['雨起','雨歇']): episode['title']=title
        for scene in data['scenes']:
            scene.update(heading='内景 空屋 白天',location='无人旧屋',time_weather='阴雨',purpose='观察风雨与积水的变化。')
            scene['beats']=[dict(beat,kind='action',speaker_id=None,text='风掀起旧报纸，雨水沿空屋的台阶流下。') for beat in scene['beats']]
        output=self.root/'silent.docx';create_docx(data,output)
        self.assertEqual(extract_docx(output),data)


if __name__=='__main__': unittest.main()
