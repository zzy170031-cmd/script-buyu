import copy
from io import BytesIO
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from fixtures import story
from screenplay_docx import create_docx, extract_docx, MARKER
from read_source_docx import read_source
from story_contract import ContractError, digest, read_json, validate_story, write_json_new, ROOT
from validate_skill import validate_skill


def mc_element(tag):
    return parse_xml(f'<mc:{tag} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"/>')


class StoryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)

    def document(self):
        path = self.base / "story.docx"
        create_docx(story(), path)
        return path

    def mutated_doc(self, action):
        path = self.document()
        doc = Document(path)
        action(doc)
        doc.save(path)
        return path

    def assert_invalid(self, mutate):
        data = story()
        mutate(data)
        with self.assertRaises(ContractError):
            validate_story(data)

    def test_skill_structure(self):
        self.assertEqual(validate_skill()["sources"], 3)

    def test_valid_contract(self):
        validate_story(story())

    def test_round_trip(self):
        self.assertEqual(digest(story()), digest(extract_docx(self.document())))

    def test_visible_dialogue_revision_is_read(self):
        def change(doc):
            paragraph = next(p for p in doc.paragraphs if p.style.name == "SBDialogue")
            # Edit a run as Word does, preserving the paragraph's stable bookmark.
            paragraph.runs[0].text = "阿砚：先把电表交给你，我来解释。"
        revised = extract_docx(self.mutated_doc(change))
        self.assertEqual(revised["scenes"][0]["beats"][1]["text"], "先把电表交给你，我来解释。")

    def test_character_table_revision(self):
        path = self.mutated_doc(lambda d: setattr(d.tables[1].cell(1, 2), "text", "修订后的性格与处境"))
        self.assertEqual(extract_docx(path)["characters"][0]["description"], "修订后的性格与处境")

    def test_multiline_unicode_and_colon(self):
        data = story()
        data["scenes"][0]["beats"][1]["text"] = "他说：灯还亮着。\n我回答：等一下。"
        path = self.base / "unicode.docx"
        create_docx(data, path)
        self.assertEqual(digest(data), digest(extract_docx(path)))

    def test_precise_decimal_budget(self):
        data = story()
        data["episodes"][0]["target_duration_seconds"] = 60.1234567
        data["project"]["target_duration_seconds"] = 120.1234567
        path = self.base / "decimal.docx"
        create_docx(data, path)
        self.assertEqual(digest(data), digest(extract_docx(path)))

    def test_unknown_paragraph_rejected(self):
        path = self.mutated_doc(lambda d: d.add_paragraph("不得漏掉的新故事情节"))
        with self.assertRaises(ContractError):
            extract_docx(path)

    def test_wrapped_paragraph_and_cell_rejected(self):
        def wrap_text(parent):
            wrapper, run, text = OxmlElement("w:customXml"), OxmlElement("w:r"), OxmlElement("w:t")
            text.text = "不能丢失的角色秘密"
            run.append(text)
            wrapper.append(run)
            parent.append(wrapper)
        for in_cell in (False, True):
            with self.subTest(in_cell=in_cell):
                path = self.base / f"wrapper-{in_cell}.docx"
                create_docx(story(), path)
                doc = Document(path)
                parent = doc.tables[1].cell(1, 2).paragraphs[0]._p if in_cell else doc.add_paragraph()._p
                wrap_text(parent)
                doc.save(path)
                with self.assertRaises(ContractError):
                    extract_docx(path)
                self.assertIn("special_text_structure_requires_review", read_source(path)["warnings"])

    def test_wrapped_table_row_warns(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "普通行"
        table.cell(1, 0).text = "特殊行的关键事实"
        wrapper = OxmlElement("w:customXml")
        row = table.rows[1]._tr
        table._tbl.remove(row)
        wrapper.append(row)
        table._tbl.append(wrapper)
        path = self.base / "wrapped-row.docx"
        doc.save(path)
        self.assertIn("special_text_structure_requires_review", read_source(path)["warnings"])

    def test_math_and_symbol_warn_and_reject(self):
        for tag in ("w:sym", "m:oMath"):
            with self.subTest(tag=tag):
                path = self.base / (tag.replace(":", "-") + ".docx")
                create_docx(story(), path)
                doc = Document(path)
                doc.paragraphs[-1].add_run()._r.append(OxmlElement(tag))
                doc.save(path)
                self.assertIn("special_text_structure_requires_review", read_source(path)["warnings"])
                with self.assertRaises(ContractError):
                    extract_docx(path)

    def test_inherited_hidden_text_rejected(self):
        def change(doc):
            hidden = doc.styles.add_style("HiddenStory", WD_STYLE_TYPE.CHARACTER)
            hidden.font.hidden = True
            child = doc.styles.add_style("DerivedStory", WD_STYLE_TYPE.CHARACTER)
            child.base_style = hidden
            doc.paragraphs[-1].runs[0].style = child
        with self.assertRaisesRegex(ContractError, "style hides"):
            extract_docx(self.mutated_doc(change))

    def test_hidden_table_style_rejected(self):
        def change(doc):
            hidden = doc.styles.add_style("HiddenCast", WD_STYLE_TYPE.TABLE)
            hidden.font.hidden = True
            doc.tables[1].style = hidden
        with self.assertRaisesRegex(ContractError, "style hides"):
            extract_docx(self.mutated_doc(change))

    def test_table_alternate_content_warns_and_rejects(self):
        def change(doc):
            table = doc.tables[1]
            branch, fallback = mc_element("AlternateContent"), mc_element("Fallback")
            row = copy.deepcopy(table.rows[1]._tr)
            # Selected branch could add a character that Table.rows never exposes.
            texts = row.xpath(".//w:t")
            texts[0].text = "C9"
            texts[1].text = "未读到的角色"
            fallback.append(row)
            branch.append(fallback)
            table._tbl.append(branch)
        path = self.mutated_doc(change)
        with self.assertRaisesRegex(ContractError, "table-level"):
            extract_docx(path)
        source = read_source(path)
        self.assertIn("alternate_content_requires_review", source["warnings"])
        self.assertIn("unsupported_table_structure_requires_review", source["warnings"])
        self.assertIn("未读到的角色", str(source["blocks"]))

    def test_row_alternate_content_rejected(self):
        def change(doc):
            branch = mc_element("AlternateContent")
            doc.tables[1].rows[1]._tr.append(branch)
        with self.assertRaisesRegex(ContractError, "row structure"):
            extract_docx(self.mutated_doc(change))

    def test_unknown_table_rejected(self):
        path = self.mutated_doc(lambda d: d.add_table(rows=1, cols=1))
        with self.assertRaises(ContractError):
            extract_docx(path)

    def test_foreign_document_rejected(self):
        path = self.mutated_doc(lambda d: setattr(d.core_properties, "identifier", "other-format"))
        with self.assertRaises(ContractError):
            extract_docx(path)

    def test_missing_bookmark_rejected(self):
        def change(doc):
            p = next(p for p in doc.paragraphs if p.style.name == "SBAction")
            p._p.remove(p._p.xpath(".//w:bookmarkStart")[0])
        with self.assertRaises(ContractError):
            extract_docx(self.mutated_doc(change))

    def test_tracked_edits_rejected_and_general_reader_warns(self):
        def change(doc):
            doc.paragraphs[-1]._p.append(OxmlElement("w:ins"))
        path = self.mutated_doc(change)
        with self.assertRaises(ContractError):
            extract_docx(path)
        self.assertIn("tracked_changes", read_source(path)["warnings"])

    def test_field_rejected(self):
        path = self.mutated_doc(lambda d: d.paragraphs[-1]._p.append(OxmlElement("w:fldSimple")))
        with self.assertRaises(ContractError):
            extract_docx(path)

    def test_image_object_rejected(self):
        path = self.mutated_doc(lambda d: d.paragraphs[-1].add_run()._r.append(OxmlElement("w:drawing")))
        with self.assertRaises(ContractError):
            extract_docx(path)
        self.assertIn("image_or_drawing", read_source(path)["warnings"])

    def test_merged_character_cells_rejected(self):
        path = self.mutated_doc(lambda d: d.tables[1].cell(1, 2).merge(d.tables[1].cell(1, 3)))
        with self.assertRaises(ContractError):
            extract_docx(path)

    def test_header_text_requires_review(self):
        path = self.mutated_doc(lambda d: setattr(d.sections[0].header.paragraphs[0], "text", "额外故事事实"))
        with self.assertRaises(ContractError):
            extract_docx(path)
        self.assertIn("额外故事事实", str(read_source(path)["side_content"]))

    def test_approval_is_not_inferred(self):
        path = self.mutated_doc(lambda d: setattr(d.tables[0].cell(2, 1), "text", "已确认"))
        with self.assertRaises(ContractError):
            extract_docx(path)

    def test_source_order(self):
        doc = Document()
        doc.add_paragraph("第一段")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "表中原文"
        doc.add_paragraph("最后一段")
        path = self.base / "source.docx"
        doc.save(path)
        result = read_source(path)
        self.assertEqual([b["type"] for b in result["blocks"]], ["p", "tbl", "p"])
        self.assertEqual(result["blocks"][1]["rows"], [["表中原文"]])
        self.assertTrue(result["requires_agent_review"])

    def test_external_links_not_followed(self):
        def change(doc):
            link = OxmlElement("w:hyperlink")
            run, text = OxmlElement("w:r"), OxmlElement("w:t")
            text.text = "外链可见文字"
            relationship = doc.part.relate_to("https://example.invalid/private", RT.HYPERLINK, is_external=True)
            link.set(qn("r:id"), relationship)
            run.append(text)
            link.append(run)
            doc.paragraphs[-1]._p.append(link)
        path = self.mutated_doc(change)
        with patch("urllib.request.urlopen", side_effect=AssertionError("unexpected network access")), patch("socket.socket", side_effect=AssertionError("unexpected network access")):
            result = read_source(path)
        self.assertIn("hyperlink_not_followed", result["warnings"])
        self.assertIn("external_relationship_not_followed", result["warnings"])
        self.assertIn("外链可见文字", str(result["blocks"]))
        with self.assertRaises(ContractError):
            extract_docx(path)

    def test_duplicate_ids(self):
        self.assert_invalid(lambda d: d["characters"].append(copy.deepcopy(d["characters"][0])))

    def test_duplicate_beats(self):
        self.assert_invalid(lambda d: d["scenes"][1]["beats"][0].update(beat_id="B1"))

    def test_duplicate_display_names(self):
        self.assert_invalid(lambda d: d["characters"][1].update(name="阿砚"))

    def test_unknown_speaker(self):
        self.assert_invalid(lambda d: d["scenes"][0]["beats"][1].update(speaker_id="C99"))

    def test_action_speaker(self):
        self.assert_invalid(lambda d: d["scenes"][0]["beats"][0].update(speaker_id="C1"))

    def test_empty_episode(self):
        self.assert_invalid(lambda d: d["scenes"].pop())

    def test_episode_order(self):
        self.assert_invalid(lambda d: d["scenes"].reverse())

    def test_budget_mismatch(self):
        self.assert_invalid(lambda d: d["project"].update(target_duration_seconds=121))

    def test_non_finite_and_bool(self):
        for value in (float("nan"), float("inf"), True):
            with self.subTest(value=value):
                self.assert_invalid(lambda d: d["project"].update(target_duration_seconds=value))

    def test_downstream_fields_rejected(self):
        for key in ("portrait_prompt", "image_prompt", "video_prompt", "shot_contracts", "model_parameters"):
            with self.subTest(key=key):
                self.assert_invalid(lambda d: d.update({key: "outside this Skill"}))

    def test_character_prompt_rejected(self):
        self.assert_invalid(lambda d: d["characters"][0].update(portrait_prompt="not permitted"))

    def test_exact_placeholder_only(self):
        self.assert_invalid(lambda d: d["scenes"][0]["beats"][0].update(text="待填写"))
        data = story()
        data["scenes"][0]["beats"][1]["text"] = "这张表还在等我填写。"
        validate_story(data)

    def test_duplicate_json_key(self):
        path = self.base / "duplicate.json"
        path.write_text('{"a":1,"a":2}', encoding="utf-8")
        with self.assertRaises(ContractError):
            read_json(path)

    def test_no_overwrite(self):
        doc = self.document()
        original = doc.read_bytes()
        with self.assertRaises(FileExistsError):
            create_docx(story(), doc)
        self.assertEqual(original, doc.read_bytes())
        path = self.base / "data.json"
        write_json_new(path, {"a": 1})
        with self.assertRaises(FileExistsError):
            write_json_new(path, {"a": 2})
        self.assertEqual(read_json(path), {"a": 1})

    def test_invalid_export_leaves_no_file(self):
        data = story()
        data["project"]["episode_count"] = 5
        target = self.base / "bad.docx"
        with self.assertRaises(ContractError):
            create_docx(data, target)
        self.assertFalse(target.exists())

    def test_dtd_rejected(self):
        for encoding in ("utf-8", "utf-16"):
            with self.subTest(encoding=encoding):
                path = self.base / (encoding + ".docx")
                xml = f'<?xml version="1.0" encoding="{encoding}"?><!DOCTYPE w:document [<!ENTITY x "harmless">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&x;</w:t></w:r></w:p></w:body></w:document>'
                with ZipFile(path, "w") as archive:
                    archive.writestr("word/document.xml", xml.encode(encoding))
                with self.assertRaisesRegex(ContractError, "DTD or entity"):
                    read_source(path)

    def test_crlf_has_actionable_error(self):
        data = story()
        data["scenes"][0]["beats"][0]["text"] = "第一行\r\n第二行"
        with self.assertRaisesRegex(ContractError, "normalize CRLF"):
            validate_story(data)

    def test_partial_write_never_publishes(self):
        def fail(stream, payload):
            stream.write(payload[:12])
            raise OSError("injected partial write")
        for suffix in ("json", "docx"):
            path = self.base / ("partial." + suffix)
            with patch("story_contract._write_stream", side_effect=fail):
                with self.assertRaisesRegex(OSError, "partial write"):
                    write_json_new(path, story()) if suffix == "json" else create_docx(story(), path)
            self.assertFalse(path.exists())
            self.assertEqual(list(self.base.glob(".*.tmp")), [])

    def test_close_failure_never_publishes(self):
        normal_open = Path.open
        class CloseFailure:
            def __init__(self, stream):
                self.stream = stream
            def __enter__(self):
                return self.stream
            def __exit__(self, *args):
                self.stream.close()
                raise OSError("injected close failure")
        def fail_close(path, *args, **kwargs):
            stream = normal_open(path, *args, **kwargs)
            return CloseFailure(stream) if path.suffix == ".tmp" else stream
        for suffix in ("json", "docx"):
            path = self.base / ("close." + suffix)
            with patch.object(Path, "open", new=fail_close):
                with self.assertRaisesRegex(OSError, "close failure"):
                    write_json_new(path, story()) if suffix == "json" else create_docx(story(), path)
            self.assertFalse(path.exists())
            self.assertEqual(list(self.base.glob(".*.tmp")), [])

    def test_publication_collision_preserves_other_file(self):
        path = self.base / "race.json"
        def collide(source, destination):
            Path(destination).write_bytes(b"other writer")
            raise FileExistsError(str(destination))
        with patch("story_contract.os.link", side_effect=collide):
            with self.assertRaises(FileExistsError):
                write_json_new(path, story())
        self.assertEqual(path.read_bytes(), b"other writer")
        self.assertEqual(list(self.base.glob(".*.tmp")), [])

    def test_cli_from_unrelated_directory(self):
        source = self.base / "story.json"
        output = self.base / "cli.docx"
        write_json_new(source, story())
        result = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/generate_story_docx.py"), "--input", str(source), "--output", str(output)], cwd=self.base, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(digest(story()), digest(extract_docx(output)))


if __name__ == "__main__":
    unittest.main()
